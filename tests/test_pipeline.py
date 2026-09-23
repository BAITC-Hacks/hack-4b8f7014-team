from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document

from app.adapters import PipelineError, align_speakers
from app.exports import export_docx, export_pdf
from app.schemas import Meeting, Minutes, MinutesReview, ProcessOptions, Segment, Task, TaskReview
from app.storage import Store
from app.worker import run_once


def test_queue_recovery_review_and_export(tmp_path):
    store = Store(tmp_path)
    meeting = Meeting(id=uuid4(), filename="test.wav", created_at=datetime.now(UTC))
    store.add_meeting(meeting)
    store.enqueue(meeting.id, ProcessOptions())
    assert store.claim().id == meeting.id
    assert store.claim() is None
    store.recover()

    class FakePipeline:
        def run(self, meeting, progress):
            progress("extracting")
            return Minutes(meeting_id=meeting.id, summary="Тестовое саммари Ә Ғ Қ Ң Ө Ұ Ү Һ І",
                           transcript=[Segment(start=0, end=2, text="Подготовить отчёт", speaker_id="S0")],
                           tasks=[Task(meeting_id=meeting.id, description="Подготовить отчёт")])

    assert run_once(store, FakePipeline())
    assert not run_once(store, FakePipeline())
    assert store.meeting(meeting.id).status == "completed"
    minutes = store.minutes(meeting.id)
    store.review_task(minutes.tasks[0].id, TaskReview(description="Подготовить отчёт", responsible="Айдана"))
    store.review_minutes(meeting.id, MinutesReview(summary="Проверено", speakers={"S0": "Айдана"}))
    minutes = store.minutes(meeting.id)
    assert minutes.tasks[0].responsible == "Айдана"
    document = Document(BytesIO(export_docx(meeting, minutes)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Айдана" in text and "Проверено" in text
    assert text.index("Текст совещания") < text.index("Саммари по ключевым пунктам") < text.index("Поручения")
    assert len(document.tables) == 2
    assert [cell.text for cell in document.tables[1].rows[0].cells] == ["Поручение", "Ответственный", "Срок"]
    with pytest.raises(PipelineError):
        export_pdf(meeting, minutes, tmp_path / "missing.ttf")
    with pytest.raises(ValueError):
        store.enqueue(meeting.id, ProcessOptions())


def test_failed_job_can_retry_without_leaking_error(tmp_path):
    store = Store(tmp_path)
    meeting = Meeting(id=uuid4(), filename="test.wav", created_at=datetime.now(UTC))
    store.add_meeting(meeting)
    store.enqueue(meeting.id, ProcessOptions())

    class BrokenPipeline:
        def run(self, meeting, progress):
            raise RuntimeError("private transcript must not escape")

    run_once(store, BrokenPipeline())
    failed = store.meeting(meeting.id)
    assert failed.status == "failed"
    assert "private" not in failed.error
    assert store.minutes(meeting.id) is None
    assert store.enqueue(meeting.id, ProcessOptions()).status == "queued"


def test_ambiguous_speakers_remain_unknown():
    result = align_speakers([Segment(start=0, end=1, text="Да")], [(0, 1, "A"), (0, 1, "B")])
    assert result[0].speaker_id is None
    result = align_speakers([Segment(start=0, end=1, text="Да")], [(0, 1, "A")])
    assert result[0].speaker_id == "A"


def test_pdf_unicode_and_markup(tmp_path):
    candidates = [Path("C:/Windows/Fonts/arial.ttf"),
                  Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    font = next((path for path in candidates if path.is_file()), None)
    if font is None:
        pytest.skip("Install a RU/KZ font for the PDF integration check")
    meeting = Meeting(id=uuid4(), filename="test.wav", created_at=datetime.now(UTC))
    minutes = Minutes(meeting_id=meeting.id, summary="Ә Ғ Қ Ң Ө Ұ Ү Һ І <tag> & текст",
                      transcript=[], tasks=[])
    content = export_pdf(meeting, minutes, font)
    assert content.startswith(b"%PDF-")
    assert b"/FontFile2" in content
