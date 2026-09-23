from hashlib import sha256
from io import BytesIO
from xml.sax.saxutils import escape

from app.adapters import PipelineError


def content_blocks(meeting, minutes):
    yield 'title', 'Протокол совещания'
    if minutes.organization:
        yield 'text', minutes.organization
    yield 'text', 'Тема: ' + (minutes.topic or meeting.filename)
    if meeting.options.meeting_date:
        yield 'text', f'Дата: {meeting.options.meeting_date}'
    yield 'note', 'Черновик ИИ. Проверьте имена, цифры, поручения и сроки.'
    yield 'heading', 'Текст совещания'
    for segment in minutes.transcript:
        name = minutes.speakers.get(segment.speaker_id, segment.speaker_id or 'Неизвестный участник')
        yield 'speaker', name
        yield 'text', segment.text.strip()
    yield 'heading', 'Саммари по ключевым пунктам'
    yield 'text', minutes.summary
    yield 'report_table', [['Направление / доклад', 'Показатель', 'Проблема']] + [
        [p.direction, p.indicator, p.problem] for p in minutes.report_points
    ] if minutes.report_points else [
        ['Направление / доклад', 'Показатель', 'Проблема'],
        ['Не заполнено', 'Не указан', 'Не указана']]
    yield 'heading', 'Поручения'
    rows = [['Поручение', 'Ответственный', 'Срок']]
    for task in minutes.tasks:
        description = task.description
        if task.evidence_status == 'unverified':
            description += '\nЦитата не подтверждена — проверить по транскрипту.'
        elif task.needs_review:
            description += '\nТребует проверки.'
        deadline = task.deadline_text or (str(task.deadline) if task.deadline else 'Не указан')
        if task.deadline and task.deadline_text:
            deadline += f' ({task.deadline})'
        rows.append([description, task.responsible or 'Не указан', deadline])
    yield 'task_table', rows if len(rows) > 1 else rows + [['Не обнаружены', '—', '—']]


def export_docx(meeting, minutes):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ['Title', 'Heading 1']:
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for kind, value in content_blocks(meeting, minutes):
        if kind.endswith('_table'):
            table = document.add_table(rows=0, cols=3)
            table.style = 'Table Grid'
            table.autofit = False
            widths = [8.0, 4.0, 4.51] if kind == 'task_table' else [5.5, 4.0, 7.01]
            for col, width in zip(table.columns, widths, strict=True):
                col.width = Cm(width)
            for index, values in enumerate(value):
                row = table.add_row()
                for cell, text, width in zip(row.cells, values, widths, strict=True):
                    cell.width = Cm(width)
                    cell.text = text
                    if index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                if index == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            document.add_paragraph()
        elif kind in {'title', 'heading'}:
            document.add_heading(value, 0 if kind == 'title' else 1)
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(value)
            if kind == 'speaker':
                run.bold = True
                paragraph.paragraph_format.keep_with_next = True
            if kind == 'note':
                run.italic = True
                run.font.size = Pt(9)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_pdf(meeting, minutes, font_path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    if not font_path.is_file():
        raise PipelineError('Для PDF укажите MINUTES_PDF_FONT: локальный TTF-шрифт с RU/KZ')
    font_name = 'Minutes' + sha256(str(font_path.resolve()).encode()).hexdigest()[:12]
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
    face = pdfmetrics.getFont(font_name).face
    if any(ord(letter) not in face.charToGlyph for letter in 'ӘәҒғҚқҢңӨөҰұҮүҺһІіЯя'):
        raise PipelineError('PDF-шрифт не содержит все необходимые RU/KZ символы')
    styles = {kind: ParagraphStyle(kind, fontName=font_name, fontSize=size,
              leading=size * 1.2, spaceAfter=6, alignment=TA_LEFT,
              keepWithNext=kind == 'speaker')
              for kind, size in [('title', 20), ('heading', 14), ('speaker', 11), ('text', 11), ('note', 9)]}
    flow = []
    def paragraph(text, kind='text'):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[kind])
    for kind, value in content_blocks(meeting, minutes):
        if kind.endswith('_table'):
            rows = [[paragraph(cell) for cell in row] for row in value]
            table = LongTable(rows, colWidths=[226, 114, 128] if kind == 'task_table' else [156, 113, 199],
                              repeatRows=1, splitByRow=1, splitInRow=1)
            table.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
            flow.extend([table, Spacer(1, 10)])
        else:
            flow.append(paragraph(value, kind))
    buffer = BytesIO()
    SimpleDocTemplate(buffer, pagesize=(612, 792), rightMargin=72, leftMargin=72,
                      topMargin=72, bottomMargin=72).build(flow)
    return buffer.getvalue()
